from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from django.conf import settings
from docx import Document
from lxml import etree

from .models import TransportUserSettings, TransportPDF, TransportRequest
from .services import build_docx, NS, W, TEMPLATE_NAME, _replace_tokens


class TransportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser(username="transport", email="transport@example.test", password="test")
        self.client.force_login(self.user)
        converter = patch('transport_app.views.generate_pdf', return_value=b'%PDF-1.4\ntest')
        self.pdf_generator = converter.start()
        self.addCleanup(converter.stop)

    def payload(self):
        data = {"organization": "Организация", "site": "Участок 5", "place": "Пл. 6", "phone": "123456", "date": "2026-10-05",
                "author": "Автор", "reviewer": "Согласующий", "approver": "Утверждающий", "version": "1",
                "items-TOTAL_FORMS": "2", "items-INITIAL_FORMS": "0"}
        for index, mode, shift in [(0, "day", "1"), (1, "period", "24")]:
            for key, value in {"vehicle": "Кран", "quantity": "2", "mode": mode, "start": "2026-10-06", "end": "2026-10-12" if index else "", "shift": shift, "frequency": "alternate", "work": "Монтаж", "note": ""}.items():
                data[f"items-{index}-{key}"] = value
        return data

    def create_request(self):
        response = self.client.post(reverse("transport:create"), self.payload())
        self.assertEqual(response.status_code, 302)
        return TransportRequest.objects.get()

    def test_mixed_schedule_and_personal_defaults(self):
        obj = self.create_request()
        single, period = list(obj.items.all())
        self.assertEqual(single.end, single.start)
        self.assertEqual(single.frequency, "daily")
        self.assertEqual(period.shift, "24")
        self.assertEqual(period.schedule, "С 06.10.2026 по 12.10.2026\nКруглосуточно\nЧерез день")
        self.assertNotIn("date", TransportUserSettings.objects.get(user=self.user).values)
        response = self.client.get(reverse("transport:create"))
        self.assertEqual(response.context["form"]["place"].value(), "Пл. 6")
        self.assertEqual(response.context["form"]["date"].value(), timezone.localdate())
        self.assertContains(self.client.get(reverse("transport:detail", args=[obj.pk])), "Круглосуточно")
        self.assertContains(self.client.get(reverse("transport:list")), "Кран")
        other = get_user_model().objects.create_superuser(username="other", email="other@example.test", password="test")
        self.client.force_login(other)
        self.assertFalse(self.client.get(reverse("transport:create")).context["form"]["place"].value())

    def test_invalid_dates_quantity_and_empty_items(self):
        for changes in [{"items-1-end": "2026-10-01"}, {"items-0-quantity": "0"}, {"items-TOTAL_FORMS": "0"}, {"items-1-shift": "99"}]:
            data = self.payload(); data.update(changes)
            self.assertEqual(self.client.post(reverse("transport:create"), data).status_code, 200)
            self.assertEqual(TransportRequest.objects.count(), 0)

    def test_edit_opens_transport_step_with_existing_data(self):
        obj = self.create_request()
        response = self.client.get(reverse('transport:edit', args=[obj.pk]))
        self.assertEqual(response.context['start_step'], 2)
        self.assertEqual(response.context['form']['place'].value(), obj.place)
        self.assertContains(response, 'data-start-step="2"')
        self.assertEqual(self.client.get(reverse('transport:create')).context['start_step'], 1)

    def test_profile_preference_skips_only_complete_defaults(self):
        url = reverse('authapp:profile')
        self.client.post(url, {'section': 'transport_app', 'transport_app-skip_general': 'on'})
        self.assertEqual(self.client.get(reverse('transport:create')).context['start_step'], 1)
        obj = self.create_request()
        self.assertTrue(obj.pdfs.exists())
        self.assertEqual(self.client.get(reverse('transport:create')).context['start_step'], 2)
        preview = self.client.get(reverse('transport:download', args=[obj.pk, 1]), {'preview': '1'})
        self.assertTrue(preview['Content-Disposition'].startswith('inline;'))
        defaults = TransportUserSettings.objects.get(user=self.user)
        self.assertTrue(defaults.skip_general)
        self.client.post(url, {'section': 'transport_app'})
        defaults.refresh_from_db()
        self.assertFalse(defaults.skip_general)
        self.assertEqual(defaults.values['place'], 'Пл. 6')
        self.assertEqual(self.client.get(reverse('transport:create')).context['start_step'], 1)
        self.assertEqual(self.client.post(url, {'section': 'unknown'}).status_code, 400)

    def test_profile_defaults_can_be_saved_without_creating_request(self):
        data = {'section': 'transport_app', 'transport_app-skip_general': 'on'}
        for name in ('organization', 'site', 'place', 'phone', 'author', 'reviewer', 'approver'):
            data['transport_app-' + name] = self.payload()[name]
        self.assertRedirects(
            self.client.post(reverse('authapp:profile'), data),
            reverse('authapp:profile') + '#application-settings',
        )
        self.assertFalse(TransportRequest.objects.exists())
        response = self.client.get(reverse('transport:create'))
        self.assertEqual(response.context['start_step'], 2)
        self.assertEqual(response.context['form']['organization'].value(), 'Организация')
        data['transport_app-place'] = ', ,'
        response = self.client.post(reverse('authapp:profile'), data)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-errors="1"')
        self.assertEqual(TransportUserSettings.objects.get(user=self.user).values['place'], 'Пл. 6')

    def test_pdf_failure_keeps_request_and_edit_discards_old_pdf(self):
        with self.assertLogs("transport_app.views", level="ERROR"), patch("transport_app.views.generate_pdf", side_effect=RuntimeError("unavailable")):
            obj = self.create_request()
        url = reverse("transport:generate", args=[obj.pk])
        obj.refresh_from_db()
        self.assertEqual(obj.revision, 1)
        self.assertFalse(obj.pdfs.exists())
        with patch("transport_app.views.generate_pdf", return_value=b"%PDF-1.4\ntest"):
            self.client.post(url, {"version": 1})
        obj.refresh_from_db()
        self.assertEqual(obj.revision, 1)
        self.assertEqual(obj.pdfs.count(), 1)
        self.client.get(reverse("transport:download", args=[obj.pk, 1]))
        obj.refresh_from_db(); self.assertEqual(obj.revision, 1)
        data = self.payload(); data["items-INITIAL_FORMS"] = "2"
        for n, item in enumerate(obj.items.all()): data[f"items-{n}-id"] = str(item.pk)
        response = self.client.post(reverse("transport:edit", args=[obj.pk]), data)
        self.assertEqual(response.status_code, 302)
        obj.refresh_from_db(); self.assertEqual(obj.revision, 2)
        self.assertEqual(obj.pdfs.get().revision, 2)
        self.assertEqual(self.client.get(reverse('transport:download', args=[obj.pk, 1])).status_code, 404)
        self.assertTrue(obj.events.filter(text='Сформирован PDF, версия 1').exists())
        with patch('transport_app.views.generate_pdf', return_value=b'%PDF-1.4\nupdated'):
            self.client.post(url, {'version': 2})
        self.assertEqual(obj.pdfs.count(), 1)
        self.assertEqual(obj.pdfs.get().revision, 2)
        self.client.post(reverse("transport:edit", args=[obj.pk]), data)
        obj.refresh_from_db(); self.assertEqual(obj.revision, 2)

    def test_list_actions_download_current_pdf_and_no_status(self):
        obj = self.create_request()
        response = self.client.get(reverse('transport:list'))
        for label in ('Дата', 'Площадка', 'Период', 'Автотранспорт', 'Действия'):
            self.assertContains(response, f'<th>{label}</th>', html=True)
        self.assertNotContains(response, 'Статус')
        self.assertNotContains(response, '<th>Версия</th>')
        self.assertNotContains(response, '<th>№ заявки</th>')
        self.assertContains(response, 'bi-file-earmark-pdf')
        self.assertContains(response, 'bi-pencil-square')
        with patch('transport_app.views.generate_pdf', return_value=b'%PDF-1.4\ntest') as generator:
            for _ in range(2):
                response = self.client.post(reverse('transport:generate', args=[obj.pk]), {'version': 1, 'action': 'download'})
                self.assertRedirects(response, reverse('transport:download', args=[obj.pk, 1]))
            generator.assert_not_called()
        response = self.client.get(reverse('transport:list'))
        self.assertContains(response, reverse('transport:download', args=[obj.pk, 1]))
        self.assertEqual(self.client.post(f'/transport/{obj.pk}/status/', {'status': 'sent'}).status_code, 404)

    def test_vehicle_search_ignores_case_spacing_and_does_not_duplicate(self):
        obj = self.create_request()
        obj.items.update(vehicle='Автокран 50 т')
        other_data = self.payload()
        other_data['items-0-vehicle'] = 'Автокран 25 т'
        other_data['items-1-vehicle'] = 'Погрузчик'
        self.client.post(reverse('transport:create'), other_data)
        for query in ('автокран 50т.', 'АВТОКРАН 50 Т', '50т'):
            response = self.client.get(reverse('transport:list'), {'q': query})
            self.assertEqual([entry.pk for entry in response.context['page']], [obj.pk])
        self.assertEqual(self.client.get(reverse('transport:list'), {'q': 'бульдозер'}).context['page'].paginator.count, 0)
        response = self.client.get(reverse('transport:list'), {'q': 'Пл. 6'})
        self.assertEqual(response.context['page'].paginator.count, 2)
        self.assertContains(response, '<th>Площадка</th>', html=True)

    def test_pdf_filename_date_and_period(self):
        from urllib.parse import unquote
        obj = self.create_request()
        url = reverse('transport:download', args=[obj.pk, 1])
        header = self.client.get(url)['Content-Disposition']
        self.assertEqual(unquote(header), "attachment; filename*=utf-8''Заявка от 05.10.2026 (на 06.10.2026-12.10.2026).pdf")
        obj.items.update(start='2026-10-08', end='2026-10-08')
        for params, disposition in [({}, 'attachment'), ({'preview': '1'}, 'inline')]:
            header = self.client.get(url, params)['Content-Disposition']
            self.assertEqual(unquote(header), disposition + "; filename*=utf-8''Заявка от 05.10.2026 (на 08.10.2026).pdf")

    def test_delete_requires_admin_and_confirmation_and_removes_related_data(self):
        obj = self.create_request()
        url = reverse('transport:delete', args=[obj.pk])
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertTrue(TransportRequest.objects.filter(pk=obj.pk).exists())
        limited = get_user_model().objects.create_user(username='deleter', email='deleter@example.test', password='test')
        limited.user_permissions.add(Permission.objects.get(codename='delete_transportrequest'))
        self.client.force_login(limited)
        self.assertEqual(self.client.post(url, {'version': 1}).status_code, 403)
        limited.is_staff = True
        limited.save()
        self.assertEqual(self.client.post(url, {'version': 99}).status_code, 302)
        self.assertTrue(TransportRequest.objects.filter(pk=obj.pk).exists())
        self.assertEqual(self.client.post(url, {'version': 1}).status_code, 302)
        self.assertFalse(TransportRequest.objects.filter(pk=obj.pk).exists())
        self.assertFalse(TransportPDF.objects.filter(request_id=obj.pk).exists())
        self.assertFalse(obj.items.exists())
        self.assertFalse(obj.events.exists())

    def test_permissions_and_post_only(self):
        obj = self.create_request()
        self.assertEqual(self.client.get(reverse("transport:generate", args=[obj.pk])).status_code, 405)
        user = get_user_model().objects.create_user(username="limited", email="limited@example.test", password="test")
        self.client.force_login(user)
        for name, args in [("list", []), ("create", []), ("detail", [obj.pk]), ("edit", [obj.pk]), ("download", [obj.pk, 1])]:
            self.assertEqual(self.client.get(reverse("transport:" + name, args=args)).status_code, 403)

    def test_document_contains_all_schedules_and_user_text(self):
        obj = self.create_request()
        with TemporaryDirectory() as directory:
            path = Path(directory) / "transport.docx"
            build_docx(obj, path)
            document = Document(path)
        self.assertEqual(len(document.tables[0].rows), 10)
        self.assertEqual(document.tables[0].cell(0, 0).text, "7567")
        self.assertEqual(document.tables[0].cell(5, 5).text, '"05" 10 2026г.')
        self.assertIn("06.10.2026\n1 смена", document.tables[0].cell(8, 2).text)
        self.assertIn("Круглосуточно", document.tables[0].cell(9, 2).text)
        self.assertEqual(document.tables[0].cell(9, 2).text, "С 06.10.2026 по 12.10.2026\nКруглосуточно\nЧерез день")
        self.assertEqual(document.tables[0].cell(2, 0).text, "Организация")
        self.assertEqual(document.tables[1].cell(0, 1).text, obj.phone)

    def test_place_token_formats_single_multiple_and_legacy_input(self):
        obj = self.create_request()
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'places.docx'
            for value, expected in [
                ('6', 'Площадка №6'),
                ('6, 1, 7', 'Площадки №6, №1, №7'),
                (' 6, ,1, 7, ', 'Площадки №6, №1, №7'),
                ('Пл. №6', 'Площадка №6'),
                ('Пл. 6', 'Площадка №6'),
                ('№6А, №01', 'Площадки №6А, №01'),
            ]:
                obj.place = value
                build_docx(obj, path)
                self.assertEqual(Document(path).tables[0].cell(7, 0).text, expected)
        data = self.payload()
        data['place'] = ', ,'
        response = self.client.post(reverse('transport:create'), data)
        self.assertIn('place', response.context['form'].errors)

    def test_template_grows_shrinks_and_preserves_formatting(self):
        obj = self.create_request()
        obj.author = 'Начальник участка, Иванов И.И.'
        obj.approver = 'Главный инженер, Петров П.П.'
        template = Path(settings.BASE_DIR) / 'document_templates/docx' / TEMPLATE_NAME
        with ZipFile(template) as package:
            original_parts = {i: package.read(i) for i in package.namelist()}
        original = etree.fromstring(original_parts['word/document.xml'])
        original_table = original.find('.//w:tbl', NS)
        prototype = original_table.findall('w:tr', NS)[8]
        def formatting(row):
            return [etree.tostring(node, method='c14n') for node in row.iter() if node.tag in {W+'trPr', W+'tcPr', W+'pPr', W+'rPr'}]
        with TemporaryDirectory() as directory:
            target = Path(directory) / 'result.docx'
            for count in (1, 11, 15):
                obj.items.all().delete()
                for n in range(count):
                    obj.items.create(vehicle=f'Кран & <{n}>', quantity=1, mode='day', start='2026-10-06', end='2026-10-06', shift='2', work='Работы')
                build_docx(obj, target)
                with ZipFile(target) as result:
                    for name, content in original_parts.items():
                        if name != 'word/document.xml': self.assertEqual(result.read(name), content, name)
                    tree = etree.fromstring(result.read('word/document.xml'))
                rows = tree.find('.//w:tbl', NS).findall('w:tr', NS)
                self.assertEqual(len(rows), count + 8)
                for row in rows[8:]: self.assertEqual(formatting(row), formatting(prototype))
                for row, source_row in zip(rows[:8], original_table.findall('w:tr', NS)[:8]):
                    self.assertEqual(formatting(row), formatting(source_row))
                text = ''.join(tree.xpath('//w:t/text()', namespaces=NS))
                self.assertNotIn('{{', text)
                self.assertIn('Иванов И.И.', text)
                self.assertIn('Главный инженер', text)
                self.assertIn('Петров П.П.', text)
                self.assertIn('Кран & <0>', text)
                self.assertIn(obj.place_display, text)
                self.assertEqual(etree.tostring(tree.find('.//w:sectPr', NS)), etree.tostring(original.find('.//w:sectPr', NS)))

    def test_split_tokens_keep_runs_and_literal_values(self):
        root = etree.fromstring(f'<w:p xmlns:w="{NS["w"]}"><w:r><w:rPr><w:b/></w:rPr><w:t>До {{{{tech</w:t></w:r><w:r><w:t>nique}}}} после {{{{count}}}}</w:t></w:r></w:p>')
        _replace_tokens(root, {'{{technique}}': 'Кран\n{{count}}', '{{count}}': 3})
        self.assertEqual(''.join(root.xpath('.//w:t/text()', namespaces=NS)), 'До Кран{{count}} после 3')
        self.assertEqual(len(root.findall('.//w:br', NS)), 1)
        self.assertIsNotNone(root.find('.//w:b', NS))

    def test_pdf_from_stale_revision_is_not_saved(self):
        obj = self.create_request()
        obj.pdfs.all().delete()
        def concurrent_edit(_):
            TransportRequest.objects.filter(pk=obj.pk).update(revision=2)
            return b"%PDF-1.4\nstale"
        with patch("transport_app.views.generate_pdf", side_effect=concurrent_edit):
            self.client.post(reverse("transport:generate", args=[obj.pk]), {"version": 1})
        self.assertFalse(TransportPDF.objects.exists())
